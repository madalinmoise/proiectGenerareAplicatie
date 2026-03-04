import customtkinter as ctk
from ui.components.base_tab import EnterpriseTab
from tkinter import scrolledtext, filedialog, messagebox
import os
from pathlib import Path
from docx import Document
import re

class WordViewerTab(EnterpriseTab):
    def setup_ui(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        # Toolbar
        toolbar = ctk.CTkFrame(self)
        toolbar.grid(row=0, column=0, sticky="ew", padx=10, pady=10)

        ctk.CTkLabel(toolbar, text="Word Document Viewer", font=("Inter", 16, "bold")).pack(side="left", padx=10)

        self.file_combo = ctk.CTkComboBox(toolbar, width=300, command=self.on_file_select)
        self.file_combo.pack(side="left", padx=10)

        ctk.CTkButton(toolbar, text="Refresh List", width=100, command=self.refresh_list).pack(side="left", padx=5)
        ctk.CTkButton(toolbar, text="Open Folder", width=100, fg_color="#475569", command=self.open_output_folder).pack(side="left", padx=5)

        # Search
        search_frame = ctk.CTkFrame(toolbar, fg_color="transparent")
        search_frame.pack(side="right", padx=10)
        self.search_var = ctk.StringVar()
        ctk.CTkEntry(search_frame, textvariable=self.search_var, placeholder_text="Search...").pack(side="left", padx=5)
        ctk.CTkButton(search_frame, text="🔍", width=40, command=self.search_text).pack(side="left")

        # Viewer Area
        self.text_area = scrolledtext.ScrolledText(self, wrap="word", font=("Segoe UI", 11), bg="#f8fafc")
        self.text_area.grid(row=1, column=0, sticky="nsew", padx=10, pady=(0, 10))
        self.text_area.configure(state="disabled")

        self.text_area.tag_config("match", background="yellow")
        self.text_area.tag_config("placeholder", foreground="#ef4444", font=("Segoe UI", 11, "bold"))

        self.refresh_list()

    def refresh_list(self):
        out_dir = self.engine.config.get('output_dir', 'output')
        if not os.path.exists(out_dir):
            self.file_combo.configure(values=[])
            return

        files = []
        for root, dirs, filenames in os.walk(out_dir):
            for f in filenames:
                if f.endswith('.docx') and not f.startswith('~$'):
                    files.append(os.path.relpath(os.path.join(root, f), out_dir))

        self.file_combo.configure(values=files)
        if files and not self.file_combo.get():
            self.file_combo.set(files[0])

    def on_file_select(self, event=None):
        rel_path = self.file_combo.get()
        if not rel_path: return
        full_path = os.path.join(self.engine.config.get('output_dir', 'output'), rel_path)
        self.load_file(full_path)

    def load_file(self, path):
        try:
            doc = Document(path)
            full_text = [p.text for p in doc.paragraphs]
            self.text_area.configure(state="normal")
            self.text_area.delete("1.0", "end")
            self.text_area.insert("1.0", "\n".join(full_text))

            # Highlight placeholders
            content = self.text_area.get("1.0", "end")
            for match in re.finditer(r'\{\{.*?\}\}|«.*?»', content):
                start = self.text_area.index(f"1.0 + {match.start()} c")
                end = self.text_area.index(f"1.0 + {match.end()} c")
                self.text_area.tag_add("placeholder", start, end)

            self.text_area.configure(state="disabled")
        except Exception as e:
            messagebox.showerror("Error", f"Could not load document: {e}")

    def search_text(self):
        term = self.search_var.get()
        if not term: return
        self.text_area.tag_remove("match", "1.0", "end")
        start_pos = "1.0"
        while True:
            start_pos = self.text_area.search(term, start_pos, stopindex="end", nocase=True)
            if not start_pos: break
            end_pos = f"{start_pos}+{len(term)}c"
            self.text_area.tag_add("match", start_pos, end_pos)
            start_pos = end_pos
        self.text_area.see("match.first")

    def open_output_folder(self):
        out_dir = self.engine.config.get('output_dir', 'output')
        if os.path.exists(out_dir):
            os.startfile(out_dir)

    def on_engine_event(self, event_type, data):
        if event_type == 'job_finished':
            self.refresh_list()
